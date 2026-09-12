# -*- coding: utf-8 -*-
"""
MODE B Excel Simulator QA driver (mirrors qa_check.py's approach for MODE A).

For each scenario: write the scenario's inputs into a fresh copy of 04_MODE_B_SIMULATOR's live
input cells, force a LibreOffice recalculation, read the recalculated RESULT cells back, and
compare against core/engine/modes/mode_b.py run on the equivalent Client Input. Also sweeps
every recalculated cell in every scenario for a raw Excel error value (#DIV/0!, #VALUE!, #N/A,
#NAME?, #REF!, #NULL!, #NUM!) — none may ever appear (UNKNOWN/ERROR must show as readable text,
never a raw Excel error).

Scenarios 14/15 exercise the "Shared Cost Status" flag (none/unresolved/invalid) — the Excel
Simulator does not implement shared-cost allocation math (see docs/features/mode_b_target_price/
SPEC.md section 9), but the flag lets it prove it never turns an unresolved/invalid shared-cost
configuration into a silently-generated number, mirroring core/engine/modes/mode_b.py's
UNSUPPORTED_SHARED_COST_ALLOCATION / INVALID_ALLOCATION_CONFIGURATION handling.

Also compares module-status TIER: Python's run_mode_b(ci)["status"] (canonical 3-tier — OK /
INCOMPLETE / ERROR, dependency_rules.md section 5) against 04_MODE_B_SIMULATOR's "Overall
Status" cell, normalized (by string prefix — the cell also carries a human-readable suffix) to
the same 3-tier vocabulary before comparing.
"""
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(HERE))

from core.engine.modes.mode_b import run_mode_b  # noqa: E402
from scenario_helpers import make_client_input_b, METRIC_KEYS_B  # noqa: E402

SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"
BASE_WB = HERE / "Pricing_Harness_Excel_Simulator_v0.5.xlsx"
SCRATCH = HERE / "_qa_scratch_mode_b"

# 04_MODE_B_SIMULATOR input/result cell rows — read directly off the built workbook rather than
# hardcoded by hand-counting, so this can never silently drift from build_workbook.py's layout.
_label_to_row = {}
_wb_for_rows = openpyxl.load_workbook(BASE_WB, read_only=True)
_ws_for_rows = _wb_for_rows["04_MODE_B_SIMULATOR"]
for _row in _ws_for_rows.iter_rows(min_row=1, max_row=30, max_col=2):
    _cell = _row[1]
    # first match wins: the same metric label is deliberately reused in the DASHBOARD CARD
    # section further down (e.g. "Required Net Sales ex VAT", "Required List Price") — the
    # RESULT section's row always comes first in reading order.
    if _cell.value and _cell.value not in _label_to_row:
        _label_to_row[_cell.value] = _cell.row
_wb_for_rows.close()

ROW_TARGET_CM = _label_to_row["Target Contribution Margin Rate (t)"]
ROW_DIRECT_COST_B = _label_to_row["Fixed-Amount Direct + Variable Cost (C)"]
ROW_NET_SALES_FEE = _label_to_row["Net-Sales Fee Rate (b)"]
ROW_GROSS_PAYMENT_FEE = _label_to_row["Gross-Payment Fee Rate (a)"]
ROW_INCLUDES_VAT_B = _label_to_row["Price Includes VAT"]
ROW_VAT_RATE_B = _label_to_row["VAT Rate (v)"]
ROW_DISCOUNT_RATE = _label_to_row["Discount Rate"]
ROW_SHARED_FLAG = _label_to_row["Shared Cost Status (Excel simulation control — not a schema field)"]

ROW_DENOMINATOR = _label_to_row["Denominator"]
ROW_REQUIRED_NET_SALES = _label_to_row["Required Net Sales ex VAT"]
ROW_REQUIRED_GROSS_PAYMENT = _label_to_row["Required Gross Payment incl VAT"]
ROW_REQUIRED_SELLING_PRICE = _label_to_row["Required Selling Price"]
ROW_REQUIRED_LIST_PRICE = _label_to_row["Required List Price"]
ROW_EXPECTED_CM = _label_to_row["Expected Contribution Margin"]
ROW_EXPECTED_CM_RATE = _label_to_row["Expected CM Rate"]
ROW_OVERALL_STATUS = _label_to_row["Overall Status"]

RESULT_ROWS = {
    "denominator": ROW_DENOMINATOR,
    "required_net_sales_ex_vat": ROW_REQUIRED_NET_SALES,
    "required_gross_payment_incl_vat": ROW_REQUIRED_GROSS_PAYMENT,
    "required_selling_price": ROW_REQUIRED_SELLING_PRICE,
    "required_list_price": ROW_REQUIRED_LIST_PRICE,
    "expected_contribution_margin": ROW_EXPECTED_CM,
    "expected_contribution_margin_rate": ROW_EXPECTED_CM_RATE,
}

ERROR_STRINGS = {"#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#REF!", "#NULL!", "#NUM!"}

# Same 15 scenarios as 05_MODE_B_PARITY_TEST in build_workbook.py (including #14/#15's
# shared-cost flag — see module docstring). shared_flag defaults to "none" for 1-13.
SCENARIOS = [
    dict(name="1. Simple — no rate cost at all (TC1)",
         target_cm_rate=0.4, direct_cost=10000, net_sales_fee_rate=0, gross_payment_fee_rate=0,
         vat_rate=0.10, includes_vat=False, discount_rate=None, shared_flag="none"),
    dict(name="2. Net-sales fee only (TC2)",
         target_cm_rate=0.3, direct_cost=10000, net_sales_fee_rate=0.1, gross_payment_fee_rate=0,
         vat_rate=0.10, includes_vat=False, discount_rate=None, shared_flag="none"),
    dict(name="3. Gross-payment fee only, VAT-exclusive display (TC3)",
         target_cm_rate=0.3, direct_cost=10000, net_sales_fee_rate=0, gross_payment_fee_rate=0.05,
         vat_rate=0.10, includes_vat=False, discount_rate=None, shared_flag="none"),
    dict(name="4. VAT-inclusive display, both fee types (TC4)",
         target_cm_rate=0.3, direct_cost=10000, net_sales_fee_rate=0.1, gross_payment_fee_rate=0.03,
         vat_rate=0.10, includes_vat=True, discount_rate=None, shared_flag="none"),
    dict(name="5. VAT-exclusive display, both fee types (TC5)",
         target_cm_rate=0.25, direct_cost=8000, net_sales_fee_rate=0.05, gross_payment_fee_rate=0.02,
         vat_rate=0.10, includes_vat=False, discount_rate=None, shared_flag="none"),
    dict(name="6. VAT-exclusive, VAT UNKNOWN, gross-payment fee (TC11)",
         target_cm_rate=0.3, direct_cost=10000, net_sales_fee_rate=0, gross_payment_fee_rate=0.05,
         vat_rate=None, includes_vat=False, discount_rate=None, shared_flag="none"),
    dict(name="7. VAT-exclusive, VAT UNKNOWN, net-sales fee only (TC12)",
         target_cm_rate=0.3, direct_cost=10000, net_sales_fee_rate=0.1, gross_payment_fee_rate=0,
         vat_rate=None, includes_vat=False, discount_rate=None, shared_flag="none"),
    dict(name="8. VAT-inclusive display, VAT UNKNOWN",
         target_cm_rate=0.3, direct_cost=10000, net_sales_fee_rate=0.1, gross_payment_fee_rate=0,
         vat_rate=None, includes_vat=True, discount_rate=None, shared_flag="none"),
    dict(name="9. Discount rate = 0",
         target_cm_rate=0.4, direct_cost=10000, net_sales_fee_rate=0, gross_payment_fee_rate=0,
         vat_rate=0.10, includes_vat=False, discount_rate=0, shared_flag="none"),
    dict(name="10. Discount rate normal value (TC6, on top of TC4)",
         target_cm_rate=0.3, direct_cost=10000, net_sales_fee_rate=0.1, gross_payment_fee_rate=0.03,
         vat_rate=0.10, includes_vat=True, discount_rate=0.10, shared_flag="none"),
    dict(name="11. Denominator <= 0 (TC7)",
         target_cm_rate=0.5, direct_cost=10000, net_sales_fee_rate=0.3, gross_payment_fee_rate=0.3,
         vat_rate=0.10, includes_vat=False, discount_rate=None, shared_flag="none"),
    dict(name="12. Explicit zero rate (TC9)",
         target_cm_rate=0.3, direct_cost=10000, net_sales_fee_rate=0, gross_payment_fee_rate=0.05,
         vat_rate=0.10, includes_vat=False, discount_rate=None, shared_flag="none"),
    dict(name="13. UNKNOWN cost (TC8)",
         target_cm_rate=0.4, direct_cost=None, net_sales_fee_rate=0, gross_payment_fee_rate=0,
         vat_rate=0.10, includes_vat=False, discount_rate=None, shared_flag="none"),
    dict(name="14. Shared cost: fixed_share (unresolved allocation)",
         target_cm_rate=0.3, direct_cost=10000, net_sales_fee_rate=0, gross_payment_fee_rate=0,
         vat_rate=0.10, includes_vat=False, discount_rate=None, shared_flag="unresolved",
         extra_cost_items=[{
             "item_id": "shared_channel_fee", "label": "공유 채널수수료", "cost_category": "variable_selling_delivery",
             "amount": None, "rate": 0.1, "currency": None, "basis": "rate_of_net_sales",
             "applies_to_component": "shared", "allocation_rule": "fixed_share",
         }]),
    dict(name="15. Shared cost: shared + direct (invalid configuration)",
         target_cm_rate=0.4, direct_cost=10000, net_sales_fee_rate=0, gross_payment_fee_rate=0,
         vat_rate=0.10, includes_vat=False, discount_rate=None, shared_flag="invalid",
         extra_cost_items=[{
             "item_id": "shared_channel_direct", "label": "모순 설정: shared + direct 수수료",
             "cost_category": "variable_selling_delivery",
             "amount": None, "rate": 0.1, "currency": None, "basis": "rate_of_net_sales",
             "applies_to_component": "shared", "allocation_rule": "direct",
         }]),
]

TOL_AMOUNT = 0.01
TOL_RATE = 0.0001
RATE_METRICS = {"expected_contribution_margin_rate"}


def _excel_overall_status_tier(text):
    """04_MODE_B_SIMULATOR's Overall Status cell reads e.g. 'ERROR — 목표 불가능 또는 입력값
    모순' — normalize by prefix to the canonical 3-tier vocabulary (dependency_rules.md
    section 5) for comparison against Python's run_mode_b(ci)["status"]."""
    if not isinstance(text, str):
        return f"UNRECOGNIZED({text!r})"
    if text.startswith("ERROR"):
        return "ERROR"
    if text.startswith("INCOMPLETE"):
        return "INCOMPLETE"
    if text.startswith("OK"):
        return "OK"
    return f"UNRECOGNIZED({text!r})"


def run_scenario(scenario):
    name = scenario["name"]
    extra_cost_items = scenario.get("extra_cost_items")
    params = {k: v for k, v in scenario.items() if k not in ("name", "extra_cost_items")}

    wb = openpyxl.load_workbook(BASE_WB)
    ws = wb["04_MODE_B_SIMULATOR"]

    def set_cell(row, value):
        # See qa_check.py's identical note: must assign .value directly to actually blank a cell.
        ws.cell(row=row, column=3).value = value

    set_cell(ROW_TARGET_CM, params["target_cm_rate"])
    set_cell(ROW_DIRECT_COST_B, params["direct_cost"])
    set_cell(ROW_NET_SALES_FEE, params["net_sales_fee_rate"])
    set_cell(ROW_GROSS_PAYMENT_FEE, params["gross_payment_fee_rate"])
    set_cell(ROW_INCLUDES_VAT_B, params["includes_vat"])
    set_cell(ROW_VAT_RATE_B, params["vat_rate"])
    set_cell(ROW_DISCOUNT_RATE, params["discount_rate"])
    set_cell(ROW_SHARED_FLAG, params["shared_flag"])

    safe_name = "".join(ch if ch.isalnum() else "_" for ch in name)[:40]
    scratch_file = SCRATCH / f"{safe_name}.xlsx"
    wb.save(scratch_file)

    subprocess.run(
        [SOFFICE, "--headless", "--convert-to", "xlsx", "--outdir", str(SCRATCH / "recalced"), str(scratch_file)],
        check=True, capture_output=True,
    )
    recalced_path = SCRATCH / "recalced" / scratch_file.name
    rwb = openpyxl.load_workbook(recalced_path, data_only=True)
    rws = rwb["04_MODE_B_SIMULATOR"]

    errors_found = []
    for sheet in rwb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value in ERROR_STRINGS:
                    errors_found.append(f"{sheet.title}!{cell.coordinate}={cell.value}")

    excel_results = {mk: rws.cell(row=RESULT_ROWS[mk], column=3).value for mk in METRIC_KEYS_B}
    excel_overall_status_raw = rws.cell(row=ROW_OVERALL_STATUS, column=3).value

    ci_params = {k: v for k, v in params.items() if k != "shared_flag"}
    ci = make_client_input_b(**ci_params, extra_cost_items=extra_cost_items)
    py_full = run_mode_b(ci)
    py_result = py_full["per_component"]["main"]
    py_module_status = py_full["status"]

    rows = []
    all_ok = True
    for mk in METRIC_KEYS_B:
        py_status = py_result[mk]["status"]
        py_value = py_result[mk]["value"]
        excel_value = excel_results[mk]
        excel_is_number = isinstance(excel_value, (int, float)) and not isinstance(excel_value, bool)

        if py_status == "OK":
            tol = TOL_RATE if mk in RATE_METRICS else TOL_AMOUNT
            ok = excel_is_number and abs(excel_value - py_value) <= tol
        else:
            # engine says UNKNOWN/ERROR -> Excel must show the SAME text, not just "not a number"
            # (an ERROR case rendered as "UNKNOWN" or vice versa would hide a real correctness bug)
            ok = (not excel_is_number) and str(excel_value) == py_status
        all_ok = all_ok and ok
        rows.append((mk, py_status, py_value, excel_value, "PASS" if ok else "FAIL"))

    excel_module_tier = _excel_overall_status_tier(excel_overall_status_raw)
    module_status_ok = excel_module_tier == py_module_status
    all_ok = all_ok and module_status_ok

    if errors_found:
        all_ok = False

    return {
        "name": name, "rows": rows, "all_ok": all_ok, "errors_found": errors_found,
        "py_module_status": py_module_status, "excel_module_tier": excel_module_tier,
        "module_status_ok": module_status_ok,
    }


def main():
    if SCRATCH.exists():
        shutil.rmtree(SCRATCH)
    SCRATCH.mkdir()
    (SCRATCH / "recalced").mkdir()

    sys.stdout.reconfigure(encoding="utf-8")
    overall_pass = True
    for scenario in SCENARIOS:
        result = run_scenario(scenario)
        overall_pass = overall_pass and result["all_ok"]
        print(f"\n=== {result['name']} ===  {'PASS' if result['all_ok'] else 'FAIL'}")
        print(f"  Module status tier: python={result['py_module_status']}  excel={result['excel_module_tier']}  "
              f"{'PASS' if result['module_status_ok'] else 'FAIL'}")
        for mk, py_status, py_value, excel_value, verdict in result["rows"]:
            print(f"  {mk:34} python={py_status:8}{py_value!s:>18}  excel={excel_value!s:>18}  {verdict}")
        if result["errors_found"]:
            print("  !! RAW EXCEL ERRORS FOUND:", result["errors_found"])

    print("\n" + "=" * 60)
    print("OVERALL:", "ALL SCENARIOS PASS" if overall_pass else "SOME SCENARIOS FAILED")
    shutil.rmtree(SCRATCH)
    return 0 if overall_pass else 1


if __name__ == "__main__":
    sys.exit(main())
