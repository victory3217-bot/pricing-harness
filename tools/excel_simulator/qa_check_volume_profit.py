# -*- coding: utf-8 -*-
"""
Volume Profit Excel Simulator QA driver (mirrors qa_check_bep.py's approach for BEP).

For each scenario: write the scenario's inputs into a fresh copy of 12_VOLUME_PROFIT_SIMULATOR's
live input cells, force a LibreOffice recalculation, read the recalculated RESULT cells back, and
compare against core/engine/modes/volume_profit.py run on the equivalent Client Input — each
metric's value, each metric's status (including ESTIMATED), the module status, and the analysis
period basis. Also sweeps every recalculated cell for a raw Excel error value (#DIV/0!, #VALUE!,
#N/A, #NAME?, #REF!, #NULL!, #NUM!) — none may ever appear, and checks the sheet's own
independent self-check cell (|MoS qty x CMu - operating profit| <= tol) never reads FAIL.

12_VOLUME_PROFIT_SIMULATOR is inherently single-component and has no Component Count / Basis
Consistency controls, so the multi-component gate and inconsistent-fixed-cost-basis cases cannot
be driven through its input cells; those are exercised in 13_VOLUME_PROFIT_PARITY_TEST, whose own
"ALL PASS" roll-up cell this driver reads for full coverage (30 cases).

Warning codes are not compared: Excel has no cell that carries them (see 00_GUIDE).
"""
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(HERE))

from core.engine.modes.volume_profit import run_volume_profit  # noqa: E402
from scenario_helpers import (  # noqa: E402
    METRIC_KEYS_VP, METRIC_LABELS_VP, make_client_input_vp, vp_reference,
)

SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"
BASE_WB = HERE / "Pricing_Harness_Excel_Simulator_v0.6.xlsx"
SCRATCH = HERE / "_qa_scratch_vp"
SHEET = "12_VOLUME_PROFIT_SIMULATOR"

INPUT_LABELS = {
    "actual_price": "Actual Price",
    "includes_vat": "Price Includes VAT",
    "vat_rate": "VAT Rate (v)",
    "direct": "Product/Service Direct Cost",
    "varfixed": "Variable Selling/Delivery Cost (Fixed Amount)",
    "haspo": "Per-Order Cost Present (Excel simulation control — not a schema field)",
    "ratenet": "Net-Sales Fee Rate (b)",
    "rategross": "Gross-Payment Fee Rate (a)",
    "fc": "Fixed Operating Cost Amount",
    "fcbasis": "Fixed Operating Cost Basis",
    "fcs": "Fixed Cost Allocation Status (Excel simulation control — not a schema field)",
    "q": "Planned Quantity per Month (units)",
    "period": "Plan Period Basis",
    "plan": "Sales Plan Present (Excel simulation control — not a schema field)",
}

_label_to_row = {}
_wb_for_rows = openpyxl.load_workbook(BASE_WB, read_only=True)
for _row in _wb_for_rows[SHEET].iter_rows(min_row=1, max_row=60, max_col=2):
    _cell = _row[1]
    if _cell.value and _cell.value not in _label_to_row:
        _label_to_row[_cell.value] = _cell.row
_wb_for_rows.close()

INPUT_ROWS = {k: _label_to_row[v] for k, v in INPUT_LABELS.items()}
RESULT_ROWS = {k: _label_to_row[METRIC_LABELS_VP[k]] for k in METRIC_KEYS_VP}
ROW_BASIS = _label_to_row["Analysis Period Basis"]
ROW_MODULE = _label_to_row["Overall Status"]
ROW_SELF_CHECK = _label_to_row["Self-Check: |MoS qty × CMu − Operating Profit| ≤ tolerance?"]

ERROR_STRINGS = {"#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#REF!", "#NULL!", "#NUM!"}
TOL_AMOUNT = 0.01
TOL_RATIO = 0.000001
RATIO_KEYS = {"operating_profit_rate", "margin_of_safety_rate"}

BASE = dict(
    actual_price=35000, includes_vat=True, vat_rate=0.10, direct=13500, varfixed=3000, haspo=False,
    ratenet=0, rategross=0.025, fc=2000000, fcbasis="per_month", fcs="component",
    q=200, period="per_month", plan="present",
)
SIMPLE = dict(actual_price=1000, includes_vat=False, vat_rate=None, direct=400, varfixed=0,
              rategross=0, fc=60000)


def sc(name, **overrides):
    scenario = dict(BASE)
    scenario.update(overrides)
    scenario["name"] = name
    return scenario


# Single-component scenarios expressible through the simulator's own input cells.
SCENARIOS = [
    sc("Worked example (SPEC section 14), per-unit costs"),
    sc("Same, per_order cost -> ESTIMATED", haspo=True),
    sc("Plan below break-even", q=100),
    sc("Plan exactly at break-even", **SIMPLE, q=100),
    sc("Plan above break-even, VAT-exclusive price", **SIMPLE, q=250),
    sc("Zero planned quantity", q=0),
    sc("Non-integer planned quantity", q=123.5),
    sc("Planned quantity blank", q=None),
    sc("Negative planned quantity", q=-5),
    sc("Period basis blank", period=None),
    sc("Unsupported plan period", period="per_year"),
    sc("CMu < 0", actual_price=1000, includes_vat=False, vat_rate=None, direct=1200, varfixed=0,
       rategross=0, fc=50000, q=10),
    sc("CMu = 0, FC > 0", actual_price=1000, includes_vat=False, vat_rate=None, direct=1000,
       varfixed=0, rategross=0, fc=50000, q=10),
    sc("CMu = 0, FC = 0", actual_price=1000, includes_vat=False, vat_rate=None, direct=1000,
       varfixed=0, rategross=0, fc=0, q=10),
    sc("No fixed-cost item at all", fc=None, fcbasis="", fcs="none"),
    sc("Fixed-cost item exists, amount blank", fc=None),
    sc("Shared fixed cost, unresolved", fc=None, fcs="unresolved"),
    sc("Shared fixed cost, blended_only", fc=None, fcs="blended_only"),
    sc("Shared + direct invalid configuration", fc=None, fcs="invalid_direct"),
    sc("Fixed-cost basis other than per_month", fcbasis="per_unit_per_month"),
    sc("Negative fixed-cost amount", fc=-10000),
    sc("Price blank", actual_price=None),
    sc("VAT-inclusive price, VAT rate blank", vat_rate=None),
    sc("Direct cost blank", direct=None),
    sc("No sales_plan -> NOT_RUN", plan="absent"),
    sc("per_order + price blank", haspo=True, actual_price=None),
    sc("per_order + CMu < 0", haspo=True, actual_price=1000, includes_vat=False, vat_rate=None,
       direct=1200, varfixed=0, rategross=0, fc=50000, q=10),
    sc("VAT-exclusive price with gross-payment fee", actual_price=32000, includes_vat=False,
       vat_rate=0.10, rategross=0.05),
    sc("Net-sales fee rate", ratenet=0.1, rategross=0, direct=200, varfixed=0),
]


def _python_reference(scenario):
    extra_items = None
    if scenario["fcs"] == "none":
        fc_value = False
    elif scenario["fcs"] == "component":
        fc_value = scenario["fc"]
    else:
        rule = {"unresolved": "by_component_revenue", "blended_only": "blended_only",
                "invalid_direct": "direct"}[scenario["fcs"]]
        fc_value = False
        extra_items = [{"item_id": "fixed_ops_shared", "label": "공유 고정운영비",
                        "cost_category": "fixed_operating_cost", "amount": None, "rate": None,
                        "currency": None, "basis": scenario["fcbasis"] or "per_month",
                        "applies_to_component": "shared", "allocation_rule": rule}]
    ci = make_client_input_vp(
        actual_price=scenario["actual_price"], includes_vat=scenario["includes_vat"],
        vat_rate=scenario["vat_rate"], direct_cost=scenario["direct"],
        variable_fixed_cost=scenario["varfixed"], per_order=scenario["haspo"],
        net_sales_fee_rate=scenario["ratenet"], gross_payment_fee_rate=scenario["rategross"],
        fixed_operating_cost=fc_value, fixed_operating_cost_basis=scenario["fcbasis"] or "per_month",
        planned_quantity=scenario["q"], period_basis=scenario["period"],
        has_sales_plan=(scenario["plan"] == "present"), extra_cost_items=extra_items,
    )
    return vp_reference(run_volume_profit(ci))


def _sweep_errors(workbook):
    found = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value in ERROR_STRINGS:
                    found.append(f"{sheet.title}!{cell.coordinate}={cell.value}")
    return found


def _recalc(workbook, filename):
    scratch_file = SCRATCH / filename
    workbook.save(scratch_file)
    subprocess.run(
        [SOFFICE, "--headless", "--convert-to", "xlsx", "--outdir", str(SCRATCH / "recalced"), str(scratch_file)],
        check=True, capture_output=True,
    )
    return openpyxl.load_workbook(SCRATCH / "recalced" / scratch_file.name, data_only=True)


def run_scenario(scenario, index):
    py_values, py_statuses, py_module, py_basis = _python_reference(scenario)

    wb = openpyxl.load_workbook(BASE_WB)
    ws = wb[SHEET]
    for key, row in INPUT_ROWS.items():
        ws.cell(row=row, column=3).value = scenario[key] if scenario[key] != "" else None

    rwb = _recalc(wb, f"vp_{index:02d}.xlsx")
    rws = rwb[SHEET]
    errors_found = _sweep_errors(rwb)

    mismatches = []
    for key in METRIC_KEYS_VP:
        excel_value = rws.cell(row=RESULT_ROWS[key], column=3).value
        excel_status = rws.cell(row=RESULT_ROWS[key], column=4).value
        py_value = py_values[key]
        tol = TOL_RATIO if key in RATIO_KEYS else TOL_AMOUNT
        if isinstance(py_value, (int, float)):
            if not isinstance(excel_value, (int, float)) or abs(excel_value - py_value) > tol:
                mismatches.append(f"{key} value: excel={excel_value!r} python={py_value!r}")
        elif excel_value != py_value:
            mismatches.append(f"{key} value: excel={excel_value!r} python={py_value!r}")
        if excel_status != py_statuses[key]:
            mismatches.append(f"{key} status: excel={excel_status!r} python={py_statuses[key]!r}")

    excel_module = rws.cell(row=ROW_MODULE, column=3).value
    if excel_module != py_module:
        mismatches.append(f"module status: excel={excel_module!r} python={py_module!r}")
    excel_basis = rws.cell(row=ROW_BASIS, column=3).value or ""
    if excel_basis != py_basis:
        mismatches.append(f"analysis_period_basis: excel={excel_basis!r} python={py_basis!r}")
    if rws.cell(row=ROW_SELF_CHECK, column=3).value == "FAIL":
        mismatches.append("SELF-CHECK FAIL: |MoS qty x CMu - operating profit| > tolerance")

    return scenario["name"], mismatches, errors_found


def main():
    if SCRATCH.exists():
        shutil.rmtree(SCRATCH)
    (SCRATCH / "recalced").mkdir(parents=True)

    total_fail = 0
    total_raw_errors = 0
    for index, scenario in enumerate(SCENARIOS, start=1):
        name, mismatches, errors_found = run_scenario(scenario, index)
        status = "PASS" if not mismatches and not errors_found else "FAIL"
        print(f"[{status}] {index:02d}. {name}")
        if mismatches:
            total_fail += 1
            for m in mismatches:
                print(f"    MISMATCH: {m}")
        if errors_found:
            total_raw_errors += len(errors_found)
            for e in errors_found:
                print(f"    RAW EXCEL ERROR: {e}")

    # 13_VOLUME_PROFIT_PARITY_TEST's own roll-up (covers the multi-component and inconsistent
    # fixed-cost-basis cases the single-component simulator sheet cannot express).
    rwb_parity = _recalc(openpyxl.load_workbook(BASE_WB), "13_parity_static.xlsx")
    rws_parity = rwb_parity["13_VOLUME_PROFIT_PARITY_TEST"]
    parity_errors = _sweep_errors(rwb_parity)
    parity_overall = None
    for row in rws_parity.iter_rows():
        for cell in row:
            if cell.value == "All values/statuses/module/basis/self-check PASS?":
                parity_overall = rws_parity.cell(row=cell.row, column=3).value
                break
        if parity_overall is not None:
            break

    print()
    print(f"{SHEET} scenarios: {len(SCENARIOS)}, Failed: {total_fail}, Raw Excel errors: {total_raw_errors}")
    print(f"13_VOLUME_PROFIT_PARITY_TEST roll-up: {parity_overall!r}, raw errors: {len(parity_errors)}")
    for e in parity_errors:
        print(f"    RAW EXCEL ERROR (13 sheet): {e}")

    if total_fail == 0 and total_raw_errors == 0 and parity_overall == "ALL PASS" and not parity_errors:
        print("VOLUME PROFIT QA: ALL PASS")
        return 0
    print("VOLUME PROFIT QA: FAILURES PRESENT")
    return 1


if __name__ == "__main__":
    sys.exit(main())
