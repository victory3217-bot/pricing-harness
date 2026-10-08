"""Build the course input worksheet (blank template, no data) for the web calculator.

Run: python tools/input_worksheet/build_input_worksheet.py
Writes docs/calculator/Pricing_Input_Worksheet.xlsx
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[2] / "docs" / "calculator" / "Pricing_Input_Worksheet.xlsx"

INPUT = PatternFill("solid", fgColor="FFF6D6")
CALC = PatternFill("solid", fgColor="E6EEE9")
HEAD = PatternFill("solid", fgColor="1E4C40")
NOTE = Font(color="666666", size=10)
THIN = Side(style="thin", color="CCCCCC")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
KINDS = ["정액 (개당)", "정액 (주문당)", "순매출 대비 %", "지불액 대비 %"]
N_DIRECT, N_VAR = 8, 8

wb = Workbook()
ws = wb.active
ws.title = "입력 정리"
for col, w in zip("ABCDE", (34, 20, 18, 22, 44)):
    ws.column_dimensions[col].width = w


def title(r, text):
    ws.cell(r, 1, text).font = Font(bold=True, size=12, color="FFFFFF")
    for c in range(1, 6):
        ws.cell(r, c).fill = HEAD


def note(r, text):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    c = ws.cell(r, 1, text)
    c.font, c.alignment = NOTE, WRAP
    ws.row_dimensions[r].height = 15 * (1 + len(text) // 70)


def cell(r, c, v=None, fill=None, fmt=None):
    x = ws.cell(r, c, v)
    x.border = BOX
    if fill:
        x.fill = fill
    if fmt:
        x.number_format = fmt
    return x


ws.cell(1, 1, "가격 진단 계산기 입력 정리 시트 / Input worksheet").font = Font(bold=True, size=14)
note(2, "노란 칸만 입력하세요. 회색 칸은 자동 계산입니다. 맨 아래 '계산기에 입력할 값'을 웹 계산기에 옮겨 적으세요. "
        "이 파일은 비어 있는 양식이며, 입력한 내용은 이 파일(또는 내려받아 연 본인의 컴퓨터)에만 저장됩니다. "
        "Yellow cells are inputs; grey cells are formulas. Copy the summary at the bottom into the calculator.")
note(3, "직접원가 = 만드는 데 드는 비용(재료비, 개당 전기·가공비, 외주 제작비 등), 변동비 = 판매·배송 비용만(수수료, 결제수수료, 배송비). "
        "같은 비용을 두 곳에 넣으면 두 번 차감됩니다. 물량과 무관한 비용(임대료 등)은 월 고정운영비에 넣습니다.")

# --- 1. basics
title(5, "1. 기본 정보 / Basics")
rows = [
    ("실제 판매가 (KRW)", "#,##0", "판매가 / Selling price"),
    ("부가세율 (%)", "0.0", "예: 10"),
    ("위 판매가는 VAT 포함가인가?", None, "예 / 아니오"),
    ("총비용을 나눌 계획 판매량 (개)", "#,##0", "개발비 등 총액 ÷ 이 수량 = 1단위당 금액 (회수 기간 전체의 계획 판매량)"),
]
for i, (lab, fmt, hint) in enumerate(rows):
    r = 6 + i
    cell(r, 1, lab)
    cell(r, 2, None, INPUT, fmt)
    ws.cell(r, 5, hint).font = NOTE
PRICE, VAT, INCVAT, QTY = "$B$6", "$B$7", "$B$8", "$B$9"
dv = DataValidation(type="list", formula1='"예,아니오"', allow_blank=True)
ws.add_data_validation(dv)
dv.add("B8")

# --- 2. direct cost
title(11, "2. 직접원가 항목 / Direct cost items (판매 1단위당으로 환산)")
for c, h in enumerate(["항목명", "입력 금액 (KRW)", "금액의 성격", "1단위당 금액 (자동)", "메모"], 1):
    x = cell(12, c, h, CALC)
    x.font = Font(bold=True)
d0 = 13
dv_basis = DataValidation(type="list", formula1='"총액,1단위당"', allow_blank=True)
ws.add_data_validation(dv_basis)
for i in range(N_DIRECT):
    r = d0 + i
    cell(r, 1, None, INPUT)
    cell(r, 2, None, INPUT, "#,##0")
    cell(r, 3, None, INPUT)
    dv_basis.add(f"C{r}")
    cell(r, 4, f'=IF(B{r}="","",IF(C{r}="","금액의 성격 선택",IF(C{r}="1단위당",B{r},IF(N({QTY})>0,B{r}/{QTY},"판매량 입력 필요"))))', CALC, "#,##0.##")
    cell(r, 5, None, INPUT)
d1 = d0 + N_DIRECT - 1
r = d1 + 1
cell(r, 1, "직접원가 합계 (1단위당)", CALC).font = Font(bold=True)
cell(r, 4, f"=SUM(D{d0}:D{d1})", CALC, "#,##0.##")
DIRECT_SUM = f"$D${r}"

# --- 3. variable
v_title = r + 2
title(v_title, "3. 변동비 항목 / Variable selling & delivery cost (판매·배송 비용만)")
for c, h in enumerate(["항목명", "유형", "입력값", "계산기에 넣을 값 (자동)", "메모"], 1):
    x = cell(v_title + 1, c, h, CALC)
    x.font = Font(bold=True)
v0 = v_title + 2
dv_kind = DataValidation(type="list", formula1='"' + ",".join(KINDS) + '"', allow_blank=True)
ws.add_data_validation(dv_kind)
for i in range(N_VAR):
    r = v0 + i
    cell(r, 1, None, INPUT)
    cell(r, 2, None, INPUT)
    dv_kind.add(f"B{r}")
    cell(r, 3, None, INPUT, "#,##0.##")
    cell(r, 4, f'=IF(C{r}="","",IF(B{r}="","유형 선택",C{r}))', CALC, "#,##0.##")
    cell(r, 5, None, INPUT)
v1 = v0 + N_VAR - 1
note(v1 + 1, "정액은 총액이 아니라 1단위당(또는 1주문당) 금액으로 적어 주세요. %는 순매출(VAT 제외 판매가) 또는 지불액(VAT 포함 결제금액) 대비 비율입니다.")

# --- 4. other
o_title = v1 + 3
title(o_title, "4. 그 밖의 값 / Other inputs")
others = [
    ("월 고정운영비 합계 (KRW)", "#,##0", "임대료·인건비·판관비 등 월 합계 (BEP·판매량 손익)"),
    ("월 계획 판매량 (개)", "#,##0", "판매량 손익에 사용"),
    ("주문당 평균 판매 수량 (개)", "0.##", "비우면 주문당 1개로 가정"),
    ("목표 공헌이익률 (%)", "0.0", "MODE B에 사용"),
    ("정가 대비 할인율 (%)", "0.0", "MODE B 정가 역산"),
    ("목표 시장가격 (KRW)", "#,##0", "MODE C에 사용"),
]
o0 = o_title + 1
for i, (lab, fmt, hint) in enumerate(others):
    cell(o0 + i, 1, lab)
    cell(o0 + i, 2, None, INPUT, fmt)
    ws.cell(o0 + i, 5, hint).font = NOTE

# --- 5. summary
s_title = o0 + len(others) + 1
title(s_title, "5. 계산기에 입력할 값 / Values to type into the calculator")
summary = [
    ("실제 판매가 (KRW)", f"=IF({PRICE}=\"\",\"\",{PRICE})", "#,##0"),
    ("부가세율 (%)", f"=IF({VAT}=\"\",\"\",{VAT})", "0.0"),
    ("VAT 포함가 체크", f'=IF({INCVAT}="","",IF({INCVAT}="예","체크","해제"))', None),
    ("직접원가 합계 (1단위당)", f"={DIRECT_SUM}", "#,##0.##"),
]
for i, (lab, f, fmt) in enumerate(summary):
    cell(s_title + 1 + i, 1, lab, CALC).font = Font(bold=True)
    cell(s_title + 1 + i, 2, f, CALC, fmt)
r = s_title + 1 + len(summary)
cell(r, 1, "변동비 항목 (유형 / 값)", CALC).font = Font(bold=True)
cell(r, 2, "계산기의 '+ 변동비 항목 추가'로 위 3번 표의 항목을 그대로 옮기세요.", CALC)
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
for i, (lab, _, _) in enumerate(others):
    rr = r + 1 + i
    cell(rr, 1, lab, CALC).font = Font(bold=True)
    cell(rr, 2, f'=IF(B{o0 + i}="","",B{o0 + i})', CALC, others[i][1])

ws.freeze_panes = "A5"
ws.sheet_view.showGridLines = False
wb.save(OUT)
print("wrote", OUT)
